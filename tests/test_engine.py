"""Integration tests use real PDFs and independent readers, never customer data."""
import contextlib
import hashlib
import http.server
import io
import json
import os
import subprocess
import sys
import tempfile
import threading
import unittest
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'engine'))
import pymupdf as fitz
from pypdf import PdfReader
from PIL import Image
from worker import execute,PDF_ACTIONS,CONVERT_ACTIONS
from common import UserError,pages


def make_fixture(path):
    doc=fitz.open()
    for idx in range(3):
        page=doc.new_page(width=595,height=842)
        page.insert_text((50,60),f'HiPDF document page {idx+1}',fontsize=24)
        page.insert_text((50,100),'SECRET-12345',fontsize=14)
        page.insert_text((50,130),'Quarterly revenue 2026: 158000',fontsize=12)
        page.insert_text((50,165),'中文测试：文档处理与数据安全',fontname='china-s',fontsize=15)
        if idx==0:
            for y in (220,255,290,325):page.draw_line((50,y),(500,y))
            for x in (50,200,350,500):page.draw_line((x,220),(x,325))
            for r,row in enumerate([['Item','Amount','Year'],['Revenue','158000','2026'],['Cost','75000','2026']]):
                for c,value in enumerate(row):page.insert_text((58+c*150,243+r*35),value,fontsize=12)
        if idx==1:
            image=Image.effect_noise((1100,750),75).convert('RGB');buf=io.BytesIO();image.save(buf,'PNG')
            page.insert_image(fitz.Rect(50,220,545,560),stream=buf.getvalue())
    doc.set_metadata({'title':'HiPDF fixture','subject':'SECRET-12345'})
    doc.embfile_add('private.txt',b'SECRET-12345')
    doc.save(path);doc.close()


class EngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory(prefix='hipdf-tests-')
        cls.root=Path(cls.temp.name);cls.source=cls.root/'fixture.pdf';make_fixture(cls.source)
        cls.original=hashlib.sha256(cls.source.read_bytes()).hexdigest()
        # Persist synthetic examples for UI checks; no user data involved.
        (ROOT/'build/examples').mkdir(parents=True,exist_ok=True)
        (ROOT/'build/examples/HiPDF-测试文档.pdf').write_bytes(cls.source.read_bytes())
    @classmethod
    def tearDownClass(cls):cls.temp.cleanup()
    def job(self,action,options=None,files=None,ai=None):
        result=execute(dict(action=action,files=[str(x) for x in (files or [self.source])],outputDir=str(self.root/'outputs'),options=options or {},ai=ai or {}))
        self.assertTrue(result['ok']);self.assertTrue(all(Path(x).is_file() for x in result['outputs']))
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(),self.original)
        return result
    def text(self,path):
        with fitz.open(path) as doc:return '\n'.join(p.get_text() for p in doc)
    def test_01_catalog_coverage(self):
        tools=json.loads((ROOT/'Resources/tools.json').read_text())
        self.assertEqual(len(tools),33)
        self.assertEqual({t['id'] for t in tools},PDF_ACTIONS|CONVERT_ACTIONS|{'htmlToPDF','summarize','translate'})
    def test_02_page_ranges(self):
        self.assertEqual(pages('3,1-2,2',3,True),[2,0,1,1])
        for invalid in ['0','4','1-9','1,,2','1;a','-1']:
            with self.assertRaises(UserError):pages(invalid,3)
    def test_03_merge_split_order(self):
        result=self.job('merge',files=[self.source,self.source]);self.assertEqual(len(PdfReader(result['outputs'][0]).pages),6)
        result=self.job('split',{'splitMode':'ranges','pages':'1-2,3'});self.assertEqual([len(PdfReader(x).pages) for x in result['outputs']],[2,1])
        result=self.job('organize',{'pages':'3,1,2,2'})
        doc=fitz.open(result['outputs'][0]);self.assertEqual(len(doc),4);self.assertIn('page 3',doc[0].get_text());doc.close()
        self.assertEqual(len(PdfReader(self.job('delete',{'pages':'2'})['outputs'][0]).pages),2)
        self.assertEqual(len(PdfReader(self.job('extract',{'pages':'2'})['outputs'][0]).pages),1)
    def test_04_rotation_crop(self):
        result=self.job('rotate',{'rotation':'90','pages':'2'});reader=PdfReader(result['outputs'][0]);self.assertEqual(reader.pages[1].rotation,90);self.assertEqual(reader.pages[0].rotation,0)
        result=self.job('crop',{'cropMargin':'20'});self.assertAlmostEqual(float(PdfReader(result['outputs'][0]).pages[0].cropbox.width),555)
    def test_05_compression(self):
        result=self.job('compress',{'compression':'balanced'});self.assertIn('158000',self.text(result['outputs'][0]));self.assertLess(Path(result['outputs'][0]).stat().st_size,self.source.stat().st_size)
    def test_06_password_roundtrip(self):
        path=self.job('encrypt',{'newPassword':'test 密码 123'})['outputs'][0]
        reader=PdfReader(path);self.assertTrue(reader.is_encrypted);self.assertEqual(reader.decrypt('wrong'),0);self.assertNotEqual(reader.decrypt('test 密码 123'),0)
        with self.assertRaises(UserError):self.job('unlock',{'password':'wrong'},[path])
        result=self.job('unlock',{'password':'test 密码 123'},[path]);self.assertFalse(PdfReader(result['outputs'][0]).is_encrypted)
    def test_07_permanent_redaction(self):
        path=self.job('redact',{'terms':'SECRET-12345'})['outputs'][0]
        self.assertNotIn('SECRET',self.text(path));reader=PdfReader(path);self.assertNotIn('SECRET',str(reader.metadata))
        with fitz.open(path) as doc:self.assertEqual(doc.embfile_count(),0)
    def test_08_watermark_numbers_edit_sign(self):
        for action,options,expected in [('watermark',{'text':'CONFIDENTIAL'},'CONFIDENTIAL'),('numbers',{'text':'Page {page} of {total}'},'Page 1 of 3'),('edit',{'text':'ADDED TEXT'},'ADDED TEXT'),('sign',{'text':'Approved by HiPDF'},'Approved by HiPDF')]:
            with self.subTest(action=action):self.assertIn(expected,self.text(self.job(action,options)['outputs'][0]))
    def test_09_forms_canonical(self):
        region={'page':0,'rect':[.1,.5,.5,.55]}
        path=self.job('forms',{'formMode':'create','fieldName':'customer','regions':[region]})['outputs'][0]
        self.assertIn('customer',PdfReader(path).get_fields())
        path=self.job('forms',{'formMode':'fill','formValues':{'customer':'Acme Ltd'}},[path])['outputs'][0]
        reader=PdfReader(path);self.assertEqual(reader.get_fields()['customer']['/V'],'Acme Ltd')
        widgets=[a.get_object() for p in reader.pages for a in p.get('/Annots',[]) if a.get_object().get('/Subtype')=='/Widget']
        self.assertTrue(widgets);self.assertTrue(widgets[0].get('/AP',{}).get('/N'))
        flattened=self.job('forms',{'formMode':'fill','formValues':{'customer':'Flat'},'flatten':True},[path])['outputs'][0]
        self.assertFalse(PdfReader(flattened).get_fields())
    def test_10_images_scan(self):
        jpgs=self.job('pdfToImage',{'dpi':'72'})['outputs'];self.assertEqual(len(jpgs),3)
        self.assertEqual(len(PdfReader(self.job('imageToPDF',files=jpgs)['outputs'][0]).pages),3)
        self.assertEqual(len(PdfReader(self.job('scan',{'scanMode':'bw'},jpgs[:1])['outputs'][0]).pages),1)
        self.assertTrue(self.job('pdfToImage',{'imageMode':'extract'})['outputs'])
    def test_11_ocr_real_scan(self):
        with fitz.open(self.source) as source:
            scan=fitz.open();p=scan.new_page();p.insert_image(p.rect,stream=source[0].get_pixmap(dpi=130).tobytes('png'))
            path=self.root/'scanned.pdf';scan.save(path);scan.close()
        result=self.job('ocr',files=[path]);text=self.text(result['outputs'][0]);self.assertIn('HiPDF',text);self.assertIn('158000',text)
    def test_12_markdown_excel(self):
        result=self.job('markdown');self.assertIn('Revenue',Path(result['outputs'][0]).read_text())
        from openpyxl import load_workbook
        book=load_workbook(self.job('pdfToExcel')['outputs'][0]);self.assertIn('Revenue',' '.join(str(c.value) for s in book for row in s for c in row))
    def test_13_editable_office_exports(self):
        from docx import Document
        from pptx import Presentation
        word=self.job('pdfToWord',{'pages':'1'})['outputs'][0];self.assertIn('HiPDF',' '.join(p.text for p in Document(word).paragraphs))
        ppt=self.job('pdfToPPT')['outputs'][0];presentation=Presentation(ppt);self.assertEqual(len(presentation.slides),3)
        self.assertIn('158000',' '.join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame))
    def test_14_office_imports_pdfa(self):
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches
        from openpyxl import Workbook
        doc=Document();doc.add_heading('HiPDF office sample',0);doc.add_paragraph('Revenue 158000');word=self.root/'sample.docx';doc.save(word)
        ppt=Presentation();slide=ppt.slides.add_slide(ppt.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(2)).text='HiPDF presentation';powerpoint=self.root/'sample.pptx';ppt.save(powerpoint)
        book=Workbook();book.active.append(['Revenue',158000]);excel=self.root/'sample.xlsx';book.save(excel)
        for action,path in [('wordToPDF',word),('pptToPDF',powerpoint),('excelToPDF',excel)]:
            with self.subTest(action=action):self.assertGreater(len(PdfReader(self.job(action,files=[path])['outputs'][0]).pages),0)
        path=self.job('pdfa')['outputs'][0]
        with fitz.open(path) as doc:self.assertIn('pdfaid:part',doc.get_xml_metadata())
    def test_15_compare_repair_workflow(self):
        changed=self.job('edit',{'text':'CHANGED'})['outputs'][0]
        result=self.job('compare',files=[self.source,changed]);self.assertEqual(len(result['outputs']),2);self.assertIn('CHANGED',Path(result['outputs'][1]).read_text())
        self.assertIn('HiPDF',self.text(self.job('repair')['outputs'][0]))
        result=self.job('workflow',{'steps':[{'tool':'extract','options':{'pages':'1-2'}},{'tool':'rotate','options':{'rotation':'180'}},{'tool':'numbers','options':{'text':'Step {page}'}}]})
        reader=PdfReader(result['outputs'][0]);self.assertEqual(len(reader.pages),2);self.assertEqual(reader.pages[0].rotation,180);self.assertIn('Step 1',self.text(result['outputs'][0]))
    def test_16_failure_cleanup(self):
        with self.assertRaises(UserError):self.job('delete',{'pages':'1-3'})
        self.assertFalse(list((self.root/'outputs').glob('.hipdf-working-*')))
    def test_17_signed_pdf_integrity(self):
        from cryptography import x509
        from cryptography.x509.oid import NameOID
        from cryptography.hazmat.primitives import hashes,serialization
        from cryptography.hazmat.primitives.asymmetric import rsa
        from cryptography.hazmat.primitives.serialization import pkcs12
        from datetime import datetime,timedelta,timezone
        from pyhanko.pdf_utils.reader import PdfFileReader
        from pyhanko.sign.validation import validate_pdf_signature
        from pyhanko_certvalidator import ValidationContext
        key=rsa.generate_private_key(public_exponent=65537,key_size=2048)
        subject=x509.Name([x509.NameAttribute(NameOID.COMMON_NAME,'HiPDF test')])
        now=datetime.now(timezone.utc)
        cert=x509.CertificateBuilder().subject_name(subject).issuer_name(subject).public_key(key.public_key()).serial_number(x509.random_serial_number()).not_valid_before(now-timedelta(days=1)).not_valid_after(now+timedelta(days=1)).sign(key,hashes.SHA256())
        path=self.root/'test.p12';path.write_bytes(pkcs12.serialize_key_and_certificates(b'hipdf',key,cert,None,serialization.BestAvailableEncryption(b'test')))
        signed=self.job('sign',{'signMode':'certificate','certificate':str(path),'certPassword':'test'})['outputs'][0]
        with open(signed,'rb') as stream:
            reader=PdfFileReader(stream);self.assertEqual(len(reader.embedded_signatures),1)
            status=validate_pdf_signature(reader.embedded_signatures[0],signer_validation_context=ValidationContext(trust_roots=[reader.embedded_signatures[0].signer_cert],allow_fetching=False))
            self.assertTrue(status.intact);self.assertTrue(status.valid)
    def test_18_ai_protocol_and_documents(self):
        from ai import endpoint,chunks,chat
        calls=[]
        class Handler(http.server.BaseHTTPRequestHandler):
            def log_message(self,*args):pass
            def do_POST(self):
                body=json.loads(self.rfile.read(int(self.headers['Content-Length'])))
                calls.append((self.path,self.headers.get('Authorization'),body))
                data=json.dumps({'choices':[{'message':{'content':'测试译文：核心结论与收入 158000。'},'finish_reason':'stop'}]},ensure_ascii=False).encode()
                self.send_response(200);self.send_header('Content-Type','application/json');self.end_headers();self.wfile.write(data)
        server=http.server.ThreadingHTTPServer(('127.0.0.1',0),Handler)
        thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
        config={'baseURL':f'http://127.0.0.1:{server.server_port}/v1','apiKey':'test-key','model':'mock','chunkSize':600}
        try:
            for action,options in [('summarize',{}),('translate',{'layout':'reflow'}),('translate',{'layout':'preserve','pages':'1'})]:
                result=self.job(action,options,ai=config);self.assertTrue(any(p.endswith('.pdf') for p in result['outputs']))
            self.assertTrue(calls);self.assertTrue(all(c[0]=='/v1/chat/completions' and c[1]=='Bearer test-key' for c in calls))
            self.assertTrue(any('page 3' in c[2]['messages'][1]['content'] for c in calls))
            long_doc=fitz.open()
            for index in range(4):
                page=long_doc.new_page()
                content=(f'Long page {index+1}. Revenue 158000. '*28)+f' LONG_END_{index+1}'
                self.assertGreaterEqual(page.insert_textbox(fitz.Rect(40,40,550,800),content,fontsize=12),0)
            long_path=self.root/'long.pdf';long_doc.save(long_path);long_doc.close()
            before=len(calls)
            self.job('summarize',files=[long_path],ai=config)
            self.assertGreater(len(calls)-before,2)
            self.assertTrue(any('LONG_END_4' in c[2]['messages'][1]['content'] for c in calls[before:]))
        finally:server.shutdown();server.server_close()
        self.assertEqual(''.join(chunks('中文测试。\n'*1000,600)),'中文测试。\n'*1000)
        self.assertEqual(endpoint('https://example.test/v1/'),'https://example.test/v1/chat/completions')
        with self.assertRaises(UserError):endpoint('http://example.test/v1')


if __name__=='__main__':unittest.main(verbosity=2)
