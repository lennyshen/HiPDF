from __future__ import annotations
import difflib
import html
import io
import json
import os
import statistics
import tempfile
from pathlib import Path
import pymupdf as fitz
from common import *


def ocr_helper():
    root = Path(__file__).resolve().parent.parent
    candidates = [os.environ.get('HIPDF_OCR', ''), str(root / 'HiPDFOCR'),
                  str(root / '.build/release/HiPDFOCR'), str(root / '.build/debug/HiPDFOCR')]
    for path in candidates:
        if path and Path(path).is_file():
            return path
    raise UserError('未找到 macOS OCR 引擎，请重新构建 HiPDF。')


def ocr_document(doc, options, directory):
    langs = str(options.get('ocrLanguages', 'zh-Hans,en-US'))
    count = 0
    with tempfile.TemporaryDirectory(prefix='hipdf-ocr-', dir=directory) as temp:
        for index in pages(options.get('pages'), len(doc)):
            page = doc[index]
            if page.get_text().strip() and not truth(options.get('forceOCR', False)):
                continue
            progress(f'离线识别第 {index+1}/{len(doc)} 页', index/len(doc))
            page.remove_rotation()
            image = Path(temp) / 'page.png'
            page.get_pixmap(dpi=200, alpha=False).save(image)
            data = json.loads(run_program([ocr_helper(), str(image), langs], timeout=180))
            if data.get('error'):
                raise UserError(data['error'])
            for item in data.get('lines', []):
                x, y, w, h = item['box']
                # Vision coordinates are bottom-left normalized; MuPDF is top-left.
                rect = fitz.Rect(x * page.rect.width, (1-y-h)*page.rect.height,
                                 (x+w)*page.rect.width, (1-y)*page.rect.height)
                rect.y1 = min(page.rect.height, rect.y1 + max(3, rect.height*.5))
                text_in_rect(page, rect, item['text'], max(4, rect.height*.65), invisible=True)
                count += 1
    return count


def pdf_to_images(doc, file, directory, o):
    outputs = []
    fmt = str(o.get('imageFormat', 'jpg'))
    dpi = int(number(o, 'dpi', 150, 36, 600))
    seen = set()
    for p in pages(o.get('pages'), len(doc)):
        page = doc[p]
        if o.get('imageMode') == 'extract':
            for index, item in enumerate(page.get_images(full=True)):
                if item[0] in seen:
                    continue
                seen.add(item[0])
                image = doc.extract_image(item[0])
                path = output_path(directory, file, f'图片_{len(seen):03}', image['ext'])
                path.write_bytes(image['image'])
                outputs.append(str(path))
        else:
            path = output_path(directory, file, f'第{p+1:03}页', fmt)
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            pix.save(str(path))
            outputs.append(str(path))
    if not outputs:
        raise UserError('文件中没有可提取的嵌入图片；可选择整页转换。')
    return outputs


def pdf_to_word(doc, file, directory, o):
    from pdf2docx import Converter
    path = output_path(directory, file, '转换', 'docx')
    with tempfile.TemporaryDirectory(dir=directory) as temp:
        source = Path(temp) / 'source.pdf'
        save_pdf(doc, source)
        converter = Converter(str(source))
        try:
            converter.convert(str(path), pages=pages(o.get('pages'), len(doc)), multi_processing=False)
        finally:
            converter.close()
    from docx import Document
    Document(path)  # validate the OOXML package
    return [str(path)], ['已重建可编辑段落和表格；复杂多栏、特殊字体及扫描件的布局可能需要校对。']


def pdf_to_ppt(doc, file, directory, o):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    presentation = Presentation()
    first = doc[0]
    presentation.slide_width = Inches(first.rect.width/72)
    presentation.slide_height = Inches(first.rect.height/72)
    for index in pages(o.get('pages'), len(doc)):
        page = doc[index]
        page.remove_rotation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        w, h = page.rect.width, page.rect.height
        sx = presentation.slide_width / w
        sy = presentation.slide_height / h
        blocks = page.get_text('dict')['blocks']
        background = fitz.open()
        background.insert_pdf(doc, from_page=index, to_page=index)
        backpage = background[0]
        if o.get('pptMode', 'editable') == 'editable':
            for block in blocks:
                if 'lines' in block:
                    backpage.add_redact_annot(block['bbox'], fill=False)
            backpage.apply_redactions(images=0, graphics=0)
        data = io.BytesIO(backpage.get_pixmap(dpi=144, alpha=False).tobytes('png'))
        slide.shapes.add_picture(data, 0, 0, width=presentation.slide_width, height=presentation.slide_height)
        background.close()
        if o.get('pptMode', 'editable') == 'editable':
            for block in blocks:
                if 'lines' not in block:
                    continue
                x0,y0,x1,y1 = block['bbox']
                shape = slide.shapes.add_textbox(int(x0*sx), int(y0*sy), max(1,int((x1-x0+6)*sx)), max(1,int((y1-y0+8)*sy)))
                frame = shape.text_frame
                frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
                frame.word_wrap = False
                for line_index, line in enumerate(block['lines']):
                    paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
                    paragraph.space_after = Pt(0)
                    for span in line['spans']:
                        run = paragraph.add_run()
                        run.text = span['text']
                        run.font.size = Pt(span['size'] * sy/12700)
                        run.font.name = span['font'].split('+')[-1]
                        run.font.bold = bool(span['flags'] & 16)
                        run.font.italic = bool(span['flags'] & 2)
                        c = span['color']
                        run.font.color.rgb = RGBColor((c>>16)&255, (c>>8)&255, c&255)
    path = output_path(directory, file, '转换', 'pptx')
    presentation.save(path)
    return [str(path)], ['可编辑模式保留图形背景并重建文字框；图片模式保持整页外观。']


def safe_cell(sheet, row, col, value):
    cell = sheet.cell(row, col, '' if value is None else str(value))
    # PDF text is data, including strings beginning with '='.
    cell.data_type = 's'


def pdf_to_excel(doc, file, directory, o):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    fallback_pages = []
    for p in pages(o.get('pages'), len(doc)):
        page = doc[p]
        tables = page.find_tables().tables
        if tables:
            for idx, table in enumerate(tables):
                sheet = workbook.create_sheet(f'第{p+1}页_表{idx+1}')
                for r, values in enumerate(table.extract(), 1):
                    for c, value in enumerate(values, 1):
                        safe_cell(sheet, r, c, value)
        else:
            sheet = workbook.create_sheet(f'第{p+1}页_文本')
            lines = page.get_text(sort=True).splitlines()
            if not lines:
                raise UserError(f'第 {p+1} 页没有可提取文本。请先运行 OCR，再转换为 Excel。')
            for r, line in enumerate(lines, 1):
                safe_cell(sheet, r, 1, line)
            fallback_pages.append(str(p+1))
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='146B62')
        for column in sheet.columns:
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(70, max(16, max(len(str(c.value or '')) for c in column)*1.5))
            for cell in column:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    path = output_path(directory, file, '转换', 'xlsx')
    workbook.save(path)
    notes = ['第 '+', '.join(fallback_pages)+' 页未识别到结构化表格，已按文本行写入独立工作表。'] if fallback_pages else []
    return [str(path)], notes


def to_markdown(doc, file, directory, o):
    lines = [f'# {Path(file).stem}', '']
    image_outputs = []
    for p in pages(o.get('pages'), len(doc)):
        page = doc[p]
        lines += [f'<!-- 第 {p+1} 页 -->', '']
        text = page.get_text('dict', sort=True)
        sizes = [s['size'] for b in text['blocks'] if 'lines' in b for l in b['lines'] for s in l['spans'] if s['text'].strip()]
        typical = statistics.median(sizes) if sizes else 12
        tables = page.find_tables().tables
        for table in tables:
            data = table.extract()
            if data:
                def row(vals):
                    return '| ' + ' | '.join(str(v or '').replace('|', '\\|').replace('\n', '<br>') for v in vals) + ' |'
                lines += [row(data[0]), row(['---']*len(data[0]))] + [row(r) for r in data[1:]] + ['']
        for block in text['blocks']:
            if 'lines' not in block or any(fitz.Rect(block['bbox']).intersects(t.bbox) for t in tables):
                continue
            for line in block['lines']:
                value = ''.join(s['text'] for s in line['spans']).strip()
                if not value:
                    continue
                size = max(s['size'] for s in line['spans'])
                prefix = '## ' if size >= typical*1.4 else '### ' if size >= typical*1.17 else ''
                lines.append(prefix + value)
            lines.append('')
        for link in page.get_links():
            uri = link.get('uri', '')
            if uri.startswith(('https://', 'http://')):
                label = page.get_textbox(link['from']).strip().replace('[', '').replace(']', '') or uri
                lines += [f'[{label}]({uri.replace(" ", "%20")})', '']
        if truth(o.get('includeImages', True)):
            for index, img in enumerate(page.get_images(full=True)):
                data = doc.extract_image(img[0])
                image_path = output_path(directory, file, f'p{p+1}_image{index+1}', data['ext'])
                image_path.write_bytes(data['image'])
                image_outputs.append(str(image_path))
                lines += [f'![第 {p+1} 页图片]({image_path.name})', '']
    if not any(page.get_text().strip() for page in doc):
        raise UserError('未找到文本层，请先运行 OCR。')
    path = output_path(directory, file, '提取', 'md')
    path.write_text('\n'.join(lines), encoding='utf-8')
    return [str(path)] + image_outputs


def compare(files, directory, o):
    if len(files) != 2:
        raise UserError('比较需要且只能选择两个 PDF 文件。')
    with open_pdf(files[0], o) as a, open_pdf(files[1], o) as b:
        left = a.get_page_text(0) if len(a) else ''
        pages_a = [p.get_text(sort=True) for p in a]
        pages_b = [p.get_text(sort=True) for p in b]
        table = difflib.HtmlDiff(wrapcolumn=80).make_table('\n'.join(pages_a).splitlines(), '\n'.join(pages_b).splitlines(),
                 html.escape(Path(files[0]).name), html.escape(Path(files[1]).name), context=True, numlines=3)
        output = fitz.open()
        changed = []
        for idx in range(max(len(a), len(b))):
            pa = a[idx] if idx < len(a) else None
            pb = b[idx] if idx < len(b) else None
            wa,ha = (pa.rect.width,pa.rect.height) if pa else (595,842)
            wb,hb = (pb.rect.width,pb.rect.height) if pb else (595,842)
            page = output.new_page(width=wa+wb+36, height=max(ha,hb)+60)
            if pa:
                page.show_pdf_page(fitz.Rect(0,40,wa,40+ha), a, idx)
            if pb:
                page.show_pdf_page(fitz.Rect(wa+36,40,wa+36+wb,40+hb), b, idx)
            text_in_rect(page, fitz.Rect(10,5,wa-5,36), f'原文件 - 第 {idx+1} 页', 12)
            text_in_rect(page, fitz.Rect(wa+46,5,wa+wb+26,36), f'对比文件 - 第 {idx+1} 页', 12)
            ta = pages_a[idx] if pa else ''
            tb = pages_b[idx] if pb else ''
            image_changed = not pa or not pb or pa.get_pixmap(matrix=fitz.Matrix(.5,.5), alpha=False).digest != pb.get_pixmap(matrix=fitz.Matrix(.5,.5), alpha=False).digest
            if ta != tb or image_changed:
                changed.append(idx+1)
            if pa and pb:
                a_words, b_words = pa.get_text('words', sort=True), pb.get_text('words', sort=True)
                matcher = difflib.SequenceMatcher(a=[w[4] for w in a_words], b=[w[4] for w in b_words], autojunk=False)
                for tag,i,j,k,l in matcher.get_opcodes():
                    if tag == 'equal':
                        continue
                    for word in a_words[i:j]:
                        page.draw_rect(fitz.Rect(word[:4]) + (0,40,0,40), color=None, fill=(1,.2,.2), fill_opacity=.23)
                    for word in b_words[k:l]:
                        page.draw_rect(fitz.Rect(word[:4]) + (wa+36,40,wa+36,40), color=None, fill=(.1,.8,.4), fill_opacity=.23)
        pdf_path = output_path(directory, files[0], '并排比较')
        save_pdf(output, pdf_path)
        report = output_path(directory, files[0], '文字差异', 'html')
        report.write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>HiPDF 文件比较</title><style>body{font:14px -apple-system,sans-serif;margin:40px;color:#20322f}table{border-collapse:collapse;width:100%}td{padding:5px;border:1px solid #ddd;vertical-align:top}.diff_add{background:#c4efcb}.diff_sub{background:#ffc8c8}.diff_chg{background:#ffe4a4}h1{font-size:26px}</style><h1>文件比较</h1><p>存在文字或视觉变化的页码：'+(', '.join(map(str,changed)) or '无')+'</p>'+table+'</html>', encoding='utf-8')
        return [str(pdf_path),str(report)], [f'共 {len(changed)} 页存在文字或视觉变化。并排 PDF 标出文字增删，HTML 提供逐行差异。']


def operate(action, files, directory, o):
    if action == 'compare':
        return compare(files, directory, o)
    outputs, notes = [], []
    for index, file in enumerate(files):
        progress(f'转换 {index+1}/{len(files)}：{Path(file).name}', index / len(files))
        if action in ('wordToPDF', 'pptToPDF', 'excelToPDF'):
            with tempfile.TemporaryDirectory(dir=directory) as temp:
                result = office_convert(file, temp, 'pdf')
                path = output_path(directory, file, '转换')
                result.replace(path)
                with open_pdf(path):
                    pass
                outputs.append(str(path))
            continue
        with open_pdf(file, o) as doc:
            if truth(o.get('useOCR', False)) and action in ('pdfToWord','pdfToPPT','pdfToExcel','markdown'):
                ocr_document(doc, o, directory)
            if action == 'pdfToImage':
                outputs += pdf_to_images(doc, file, directory, o)
            elif action in ('pdfToWord', 'pdfToPPT', 'pdfToExcel'):
                paths, messages = {'pdfToWord':pdf_to_word,'pdfToPPT':pdf_to_ppt,'pdfToExcel':pdf_to_excel}[action](doc,file,directory,o)
                outputs += paths
                notes += messages
            elif action == 'ocr':
                count = ocr_document(doc, o, directory)
                if count == 0 and not any(p.get_text().strip() for p in doc):
                    raise UserError('未识别到文字，请检查扫描质量或调整识别语言。')
                outputs.append(save_pdf(doc, output_path(directory,file,'OCR')))
                notes.append(f'新增 {count} 行可搜索文字；已有文本的页面默认跳过。')
            elif action == 'markdown':
                outputs += to_markdown(doc,file,directory,o)
            elif action == 'pdfa':
                with tempfile.TemporaryDirectory(dir=directory) as temp:
                    source = Path(temp)/'source.pdf'
                    save_pdf(doc,source)
                    out_dir = Path(temp)/'out'
                    out_dir.mkdir()
                    result = office_convert(source,out_dir,'pdf:draw_pdf_Export:{"SelectPdfVersion":{"type":"long","value":"2"}}','draw_pdf_import')
                    with fitz.open(result) as check:
                        if 'pdfaid:part' not in check.get_xml_metadata():
                            raise UserError('归档引擎未生成 PDF/A 标记。')
                    path = output_path(directory,file,'PDF-A-2b')
                    result.replace(path)
                    outputs.append(str(path))
                notes.append('按 PDF/A-2b 导出；正式归档前建议使用 veraPDF 对具体输出做独立合规校验。')
            else:
                raise UserError('未知转换操作：'+action)
    return outputs,notes
